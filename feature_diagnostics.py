#!/usr/bin/env python3
"""
Targeted diagnostics for cover letter generation, resume tailoring, and job matching.
"""

import os
import sys
import tempfile
from pathlib import Path

from docx import Document

from cover_letter_generator import CoverLetterGenerator
from matcher.job_matcher import JobMatcher
from resume_tailor import ResumeTailor


PLACEHOLDERS = (
    "[your interest in the company/role - please customize this section]",
    "[your name]",
)


def print_result(status: str, name: str, details: str) -> None:
    symbols = {
        "PASS": "✓",
        "WARN": "⚠",
        "FAIL": "✗",
    }
    print(f"{symbols[status]} {status:<4} {name}")
    print(f"    {details}")


def check_job_matcher() -> tuple[str, str]:
    profile = {
        "skills": ["Python", "SQL", "Flask", "Selenium"],
        "experience": [],
        "education": [],
    }
    matcher = JobMatcher(profile)

    good_match = {
        "title": "Software Engineering Intern",
        "description": (
            "Seeking a Python and SQL intern to build backend tools with Flask "
            "for students interested in software engineering."
        ),
        "location": "Toronto, ON",
    }
    senior_role = {
        "title": "Senior Software Engineer",
        "description": "Python platform role requiring 5+ years of experience.",
        "location": "Toronto, ON",
    }
    unrelated_role = {
        "title": "Marketing Intern",
        "description": "Social media and brand marketing internship.",
        "location": "Toronto, ON",
    }

    good_score = matcher.calculate_match_score(good_match)
    senior_score = matcher.calculate_match_score(senior_role)
    unrelated_score = matcher.calculate_match_score(unrelated_role)

    if good_score <= senior_score or good_score <= unrelated_score:
        return (
            "FAIL",
            f"Expected the relevant internship to outscore comparison jobs, got "
            f"{good_score=} {senior_score=} {unrelated_score=}.",
        )

    return (
        "PASS",
        f"Relevant internship scored {good_score}/100 versus senior role {senior_score}/100 "
        f"and unrelated internship {unrelated_score}/100.",
    )


def check_cover_letter() -> tuple[str, str]:
    generator = CoverLetterGenerator()
    user_profile = {
        "skills": ["Python", "SQL", "Flask", "Selenium"],
        "experience": [{"title": "Software Developer Intern", "company": "Acme"}],
        "education": [{"degree": "BSc", "field": "Computer Science", "institution": "York University"}],
        "contact_info": {"first_name": "Mustafa Alp", "last_name": "Ari"},
    }
    job = {
        "title": "Software Engineering Intern",
        "company": "ExampleCo",
        "location": "Toronto, ON",
        "description": "Looking for a Python and SQL developer intern to build web tools with Flask.",
    }

    with tempfile.TemporaryDirectory(prefix="jobapplier_cover_letter_") as tmpdir:
        text, path = generator.generate_cover_letter(job, user_profile, output_dir=tmpdir)

        if not os.path.exists(path):
            return "FAIL", "Generator returned a path, but no file was written."

        lowered = text.lower()
        leaked_placeholders = [placeholder for placeholder in PLACEHOLDERS if placeholder in lowered]
        if leaked_placeholders:
            return "FAIL", f"Template placeholders leaked into output: {', '.join(leaked_placeholders)}."

        notes = []
        if "Mustafa Alp Ari" not in text:
            notes.append("candidate signature was not populated from profile data")
        if not generator.use_ai:
            notes.append("OpenAI key not configured, so output is template-based rather than AI-generated")

        if notes:
            return "WARN", f"Cover letter file was generated successfully at {Path(path).name}; " + "; ".join(notes) + "."

        return "PASS", f"Cover letter file was generated successfully at {Path(path).name} with no placeholder text."


def check_resume_tailoring() -> tuple[str, str]:
    tailor = ResumeTailor()
    user_profile = {
        "skills": ["Python", "Flask", "SQL", "Selenium", "JavaScript"],
        "experience": [{"title": "Builder", "company": "Acme"}],
        "education": [{"degree": "BSc", "field": "Computer Science", "institution": "York University"}],
        "contact_info": {"first_name": "Mustafa Alp", "last_name": "Ari"},
    }
    job = {
        "title": "Backend Engineering Intern",
        "company": "ExampleCo",
        "location": "Remote",
        "description": "Seeking an intern with Python, Flask, SQL, API development, and backend engineering experience.",
    }

    with tempfile.TemporaryDirectory(prefix="jobapplier_resume_tailor_") as tmpdir:
        resume_path = os.path.join(tmpdir, "resume.docx")
        doc = Document()
        doc.add_paragraph("Mustafa Alp Ari")
        doc.add_paragraph("Professional Summary")
        doc.add_paragraph("Computer science student building Python automation tools and web apps.")
        doc.add_paragraph("Skills")
        doc.add_paragraph("Python, Flask, SQL, Selenium, JavaScript")
        doc.add_paragraph("Experience")
        doc.add_paragraph("Built automation scripts and dashboards for application tracking.")
        doc.save(resume_path)

        output_path = tailor.tailor_resume(resume_path, job, user_profile, output_dir=tmpdir)
        if not os.path.exists(output_path):
            return "FAIL", "Tailor returned a path, but no resume file was written."

        out_doc = Document(output_path)
        output_text = "\n".join(paragraph.text for paragraph in out_doc.paragraphs)
        if "ExampleCo" not in output_text or "Backend Engineering Intern" not in output_text:
            return "FAIL", "Tailored resume did not include the target job metadata."

        if not tailor.use_ai:
            return (
                "WARN",
                "Tailored resume file was generated successfully, but OpenAI is not configured, "
                "so tailoring is keyword-based rather than semantic rewriting.",
            )

        return "PASS", "Tailored resume file was generated successfully with AI-assisted semantic rewriting enabled."


def main() -> int:
    print("=" * 70)
    print("JobApplier Feature Diagnostics")
    print("=" * 70)

    checks = [
        ("Job Matching", check_job_matcher),
        ("Cover Letter Generation", check_cover_letter),
        ("Resume Tailoring", check_resume_tailoring),
    ]

    statuses = []
    for name, check in checks:
        status, details = check()
        statuses.append(status)
        print_result(status, name, details)

    print("\n" + "=" * 70)
    fail_count = statuses.count("FAIL")
    warn_count = statuses.count("WARN")
    pass_count = statuses.count("PASS")
    print(f"Summary: {pass_count} pass, {warn_count} warning, {fail_count} fail")

    if fail_count:
        print("Result: feature checks found at least one reproducible failure.")
        return 1

    if warn_count:
        print("Result: core flows work, but some features are running in fallback mode.")
        return 0

    print("Result: all targeted feature diagnostics passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
